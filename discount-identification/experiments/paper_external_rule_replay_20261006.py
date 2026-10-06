"""Audit reviewed external GP policy methods on real archived price paths.

Core methods are extracted unchanged from a fixed public repository commit.
Only a small mock market supplies target persistence and constant variance.
Actual market prices drive counterfactual exposures. Actions are program
replays, not observed institutional trades or measured trading costs.
"""
from pathlib import Path
import ast,json,csv,hashlib,types,datetime,warnings
from fractions import Fraction as F
from functools import lru_cache
import numpy as np
import openpyxl
from paper_self_calibration_20261006 import closed_recover,self_calibration_certificate,pencil
from paper_calibration_audit_20261006 import point
from independent_tracking_controller_20261006 import solve_policy

ROOT=Path(__file__).resolve().parent/'paper_breakthrough_20261006'
EXT=ROOT/'external'
HORIZONS=[5,20,60]
NODES=[1-2/(h+1) for h in HORIZONS]
WEIGHTS=np.array([1/3]*3)
FILTER=np.array([.5,.35,.15])

def reviewed_external_class():
    with warnings.catch_warnings():
        warnings.simplefilter('ignore',SyntaxWarning)
        tree=ast.parse((EXT/'agents.py').read_text())
    selected=[]
    for node in tree.body:
        if isinstance(node,ast.ClassDef):
            selected += [n for n in node.body if isinstance(n,ast.FunctionDef) and n.name in
                         ['_get_gp_trade','_get_a','_get_gp_rescaling','_update_trade_by_strategyType']]
    assert len(selected)==4
    module=ast.Module(body=[ast.ClassDef(name='ReviewedGP',bases=[],keywords=[],body=selected,decorator_list=[])],type_ignores=[])
    namespace={'np':np,'RiskDriverType':types.SimpleNamespace(PnL='PnL'),
               'StrategyType':types.SimpleNamespace(Unconstrained='Unconstrained',LongOnly='LongOnly')}
    exec(compile(ast.fix_missing_locations(module),str(EXT/'agents.py'),'exec'),namespace)
    return namespace['ReviewedGP']

GP=reviewed_external_class()

def make_agent(alpha,beta,cost,s):
    agent=GP();agent.gamma=beta;agent.kappa=alpha/beta;agent.lam=cost
    agent.strategyType='Unconstrained';agent.market=types.SimpleNamespace(riskDriverType='PnL')
    agent._read_Phi=lambda:1-s
    return agent

def step(agent,p,d):
    # In the public implementation target=pnl/(risk_aversion*variance).
    return p+agent._get_gp_trade(p,agent.kappa*d,1.)

@lru_cache(maxsize=200)
def input_forecast_gain(alpha,beta,cost,s,r):
    A=np.array([[0.,0.,0.],[0.,s,(1-s)*r],[0.,0.,r]])
    B=np.array([1.,0.,0.]);Q=np.diag([cost,alpha,0.]);N=np.array([-cost,-alpha,0.]);R=alpha+cost
    P=np.zeros((3,3))
    for iteration in range(10000):
        v=N+beta*A.T@P@B;den=R+beta*B@P@B
        Pn=Q+beta*A.T@P@A-np.outer(v,v)/den
        if np.max(abs(Pn-P))<2e-14*max(1.,np.max(abs(Pn))):P=Pn;break
        P=Pn
    else:raise ArithmeticError('Input forecast Bellman solve failed')
    gain=-(N+beta*A.T@P@B)/(R+beta*B@P@B)
    core,_i,_r=solve_policy(alpha,beta,cost,s)
    assert np.max(abs(gain[:2]-core))<1e-11
    return float(gain[2])

def paired_prefix(alpha,beta,costs,L=32,report_quantum=1e-12,baseline=None,forecast_r=0.,forecast_intercept=0.):
    pairs=[];unrounded=[]
    for cost in costs:
        traces=[]
        for sign in [-1.,1.]:
            y=np.zeros(L)
            for j,s in enumerate(NODES):
                agent=make_agent(alpha,beta,cost,s);p=-.2*(j+1);d=.3*(j+1)
                feed=input_forecast_gain(alpha,beta,cost,s,forecast_r)
                E=step(agent,1.,0.);C=step(agent,0.,1.);bE=beta*E
                mean=forecast_intercept/(1-forecast_r)
                intercept_feed=C*(1-s)*mean*(bE/(1-bE)-bE*forecast_r/(1-bE*forecast_r))
                for n in range(L):
                    u=0. if baseline is None else baseline[n]
                    d=s*d+(1-s)*(u+(sign if n==0 else 0.))
                    p=step(agent,p,d)+feed*u+intercept_feed;y[n]+=WEIGHTS[j]*p
            traces.append(np.convolve(y,FILTER)[:L])
        unrounded.append((traces[1]-traces[0])/2)
        pairs.append(np.round(np.array(traces)/report_quantum)*report_quantum)
    raw=np.array(pairs);z=(raw[:,1]-raw[:,0])/2
    # Conservative deterministic enclosure: rounding plus float guard.
    t=report_quantum+1e-13
    assert np.max(abs(z-np.array(unrounded)))<=t
    # Exact rational bisection independently encloses the ideal response.
    # These finite validations establish that the declared floating guard is
    # ample for these replay records, rather than assuming it from a point fit.
    bound=exact_kernel_error(alpha,beta,costs,z)
    assert bound<=t
    return z,raw,t

def exact_kernel_error(alpha,beta,costs,observed):
    def plus(x,y):return x[0]+y[0],x[1]+y[1]
    def times(x,y):
        vals=[a*b for a in x for b in y];return min(vals),max(vals)
    maximum=F(0);L=observed.shape[1]
    for i,cost in enumerate(costs):
        A,B,K=F(alpha),F(beta),F(cost);left,right=F(0),F(1)
        for _ in range(140):
            middle=(left+right)/2
            value=B*K*middle**2-(A+K*(1+B))*middle+K
            if value>0:left=middle
            else:right=middle
        E=(left,right);aggregate=[(F(0),F(0)) for _ in range(L)]
        for j,s0 in enumerate(NODES):
            s=F(s0);num=times((1-right,1-left),(1-B*right,1-B*left))
            denom=(1-B*right*s,1-B*left*s);C=(num[0]/denom[1],num[1]/denom[0])
            p=(F(0),F(0));d=F(0)
            for n in range(L):
                d=s*d+(1-s)*(1 if n==0 else 0)
                p=plus(times(E,p),times(C,(d,d)))
                aggregate[n]=plus(aggregate[n],times((F(WEIGHTS[j]),F(WEIGHTS[j])),p))
        for n in range(L):
            y=(F(0),F(0))
            for j in range(min(n+1,len(FILTER))):y=plus(y,times((F(FILTER[j]),F(FILTER[j])),aggregate[n-j]))
            value=F(observed[i,n]);maximum=max(maximum,abs(value-y[0]),abs(value-y[1]))
    return float(maximum)

def prices_and_configuration():
    book=openpyxl.load_workbook(EXT/'commodities-summary-statistics.xlsx',read_only=True,data_only=True)
    rows=list(book['Simplified contract multiplier'].values);header=rows[0]
    config={r[0]:dict(zip(header,r)) for r in rows[1:] if r[0] and r[-1] is not None}
    book=openpyxl.load_workbook(EXT/'assets_data.xlsx',read_only=True,data_only=True)
    out={}
    for ticker in sorted(config):
        if ticker not in book.sheetnames:continue
        rows=list(book[ticker].values)[1:]
        rows=[(r[0],float(r[1])) for r in rows if r[0] is not None and isinstance(r[1],(int,float)) and np.isfinite(r[1])]
        rows.sort(key=lambda r:r[0]);assert len({r[0] for r in rows})==len(rows)
        out[ticker]=(rows,config[ticker])
    return out

def official_prices():
    manifest=json.loads((EXT/'eia_source_manifest.json').read_text())
    result={}
    for record in manifest['sources']:
        file=EXT/(record['name']+'.csv');raw=EXT/(record['name']+'.xls')
        assert hashlib.sha256(file.read_bytes()).hexdigest()==record['csv_sha256']
        assert hashlib.sha256(raw.read_bytes()).hexdigest()==record['raw_sha256']
        rows=[(datetime.datetime.fromisoformat(r['date']),float(r['price'])) for r in csv.DictReader(file.open())]
        result[record['name']]=(rows,record)
    return result

def task_paths(prices,alpha,cost,beta,forecast=False):
    changes=np.diff(np.array([r[1] for r in prices]));split=int(.7*len(changes))
    sigma=float(np.std(changes[:split],ddof=1));center=float(np.mean(changes[:split]))
    assert sigma>0
    u=(changes-center)/sigma
    r=float(np.dot(u[1:split],u[:split-1])/np.dot(u[:split-1],u[:split-1]))
    r=float(np.clip(r,-.95,.95))
    goals=np.zeros((len(u),3));positions=np.zeros_like(goals)
    private_losses=np.zeros_like(goals)
    for j,s in enumerate(NODES):
        agent=make_agent(alpha,beta,cost,s);p=d=0.
        feed=input_forecast_gain(alpha,beta,cost,s,r) if forecast else 0.
        for n,signal in enumerate(u):
            d=s*d+(1-s)*signal;pn=step(agent,p,d)+feed*signal
            private_losses[n,j]=.5*alpha*(pn-d)**2+.5*cost*(pn-p)**2
            goals[n,j]=d;positions[n,j]=pn;p=pn
    loss=private_losses@WEIGHTS
    meta=dict(train_days=split,test_days=len(changes)-split,
        train_scale=sigma,train_center=center,start=prices[0][0].isoformat(),end=prices[-1][0].isoformat(),
        test_start=prices[split+1][0].isoformat(),negative_price_days=sum(r[1]<0 for r in prices),
        training_ar1=r,forecast_module=forecast)
    if isinstance(prices[0][0],datetime.datetime):
        years=np.array([r[0].year for r in prices[split+1:]])
        meta['year_mean_losses']={str(y):float(loss[split:][years==y].mean()) for y in sorted(set(years))}
    return loss[split:],positions[split:],goals[split:],meta

def main():
    gamma=dict(list(csv.reader((EXT/'settings.csv').open()))[1:])['gamma'];task_beta=float(gamma)
    market=prices_and_configuration();asset_results=[];logs=[];audits=[]
    # Risk/cost ratios come from the external repository, not fitted to outcomes.
    # All configured matched assets are retained; no profitability selection.
    for ticker,(prices,configuration) in market.items():
        ratio=float(configuration['kappa'])/float(configuration['lam'])
        alpha=task_beta*ratio;cost=1.
        cases={}
        for beta in [.45,.6,.8,.95,.99,task_beta]:
            losses,positions,goals,meta=task_paths(prices,alpha,cost,beta)
            cases[str(beta)]=dict(mean_loss=float(losses.mean()),median_loss=float(np.median(losses)),
                                 worst_loss=float(losses.max()),observations=len(losses))
        base=cases[str(task_beta)]['mean_loss']
        asset_results.append(dict(ticker=ticker,external_risk_to_cost_ratio=ratio,alpha=alpha,
            configuration=configuration,chronological_split=meta,controller_cases=cases,
            improvement_vs_045=1-base/cases['0.45']['mean_loss'],improvement_vs_099=1-base/cases['0.99']['mean_loss']))
    # The external ratios happen to be common; use the first prescribed asset.
    alpha=asset_results[0]['alpha']
    primary_results=[]
    for name,(prices,source) in official_prices().items():
        cases={};year_losses={};forecast_cases={}
        for beta in [.45,.6,.8,.95,.99,task_beta]:
            losses,positions,goals,meta=task_paths(prices,alpha,1.,beta)
            cases[str(beta)]=dict(mean_loss=float(losses.mean()),observations=len(losses))
            year_losses[str(beta)]=meta['year_mean_losses']
            forecast_loss,_p,_d,forecast_meta=task_paths(prices,alpha,1.,beta,forecast=True)
            forecast_cases[str(beta)]=dict(mean_loss=float(forecast_loss.mean()),observations=len(forecast_loss),
                                          year_mean_losses=forecast_meta['year_mean_losses'])
        primary_results.append(dict(series=name,source=source,chronological_split=meta,controller_cases=cases,
            year_losses=year_losses,forecast_cases=forecast_cases,forecast_training_ar1=forecast_meta['training_ar1'],
            forecast_improvement_vs_045=1-forecast_cases[str(task_beta)]['mean_loss']/forecast_cases['0.45']['mean_loss'],
            forecast_improvement_vs_080=1-forecast_cases[str(task_beta)]['mean_loss']/forecast_cases['0.8']['mean_loss'],
            improvement_from_forecast=1-forecast_cases[str(task_beta)]['mean_loss']/cases[str(task_beta)]['mean_loss'],
            improvement_vs_045=1-cases[str(task_beta)]['mean_loss']/cases['0.45']['mean_loss'],
            improvement_vs_080=1-cases[str(task_beta)]['mean_loss']/cases['0.8']['mean_loss'],
            improvement_vs_099=1-cases[str(task_beta)]['mean_loss']/cases['0.99']['mean_loss']))
    rows=market['gasoil'][0];vals=np.array([r[1] for r in rows]);delta=np.diff(vals);j=int(np.argmax(abs(delta)))
    data_quality=dict(gasoil_largest_jump=dict(before_date=rows[j][0].isoformat(),after_date=rows[j+1][0].isoformat(),
        before_value=float(vals[j]),after_value=float(vals[j+1]),ratio=float(vals[j]/vals[j+1])),
        interpretation='Approximately 100-fold price-scale discontinuity; original results retained, not silently repaired.')
    prices=[(datetime.datetime.fromisoformat(r['date']),float(r['price'])) for r in csv.DictReader((EXT/'WTI_EIA_full.csv').open())]
    changes=np.diff([p for _date,p in prices]);split=int(.7*len(changes))
    split=sum(day<=datetime.datetime(2022,5,23) for day,p in prices)-1
    baseline=(changes[split:split+32]-changes[:split].mean())/changes[:split].std(ddof=1)
    training_u=(changes[:split]-changes[:split].mean())/changes[:split].std(ddof=1)
    forecast_c,forecast_r=np.linalg.lstsq(np.stack([np.ones(split-1),training_u[:-1]],1),training_u[1:],rcond=None)[0]
    forecast_r=float(np.clip(forecast_r,-.95,.95));forecast_c=float(forecast_c)
    for beta in [task_beta,.8]:
        z,raw,t=paired_prefix(alpha,beta,[1.,2.,4.],baseline=baseline,forecast_r=forecast_r,forecast_intercept=forecast_c)
        rec,sv,w=closed_recover(z);ci=self_calibration_certificate(z,t)
        exact,_,_=closed_recover(paired_prefix(alpha,beta,[1.,2.,4.],report_quantum=1e-15,baseline=baseline,
            forecast_r=forecast_r,forecast_intercept=forecast_c)[0])
        assert np.max(abs(exact-[alpha,beta,2.,4.]))<1e-3
        assert ci['beta_interval'][0]<=beta<=ci['beta_interval'][1]
        lo,hi=ci['beta_interval'];cut=.99
        decision='accept' if lo>=cut else 'fail_clock' if hi<cut else 'inconclusive'
        audits.append(dict(controller_beta=beta,alpha=alpha,reported_costs_not_used=True,
            recovered=rec.tolist(),pencil_singular_values=sv.tolist(),certificate=ci,
            clock_cut=cut,action=decision,nominal_cost_error_witness=point(z,[1.,2.02,4.]).tolist(),
            exact_rational_response_error=exact_kernel_error(alpha,beta,[1.,2.,4.],z)))
        logs.append(dict(replay_label='configured_clock' if beta==task_beta else 'short_clock',
            paired_reported_aggregate_positions=raw.tolist(),report_quantum=1e-12,response_error_bound=t,
            lags=32,input_amplitude=1.,cost_labels=['A','B','C'],
            independent_baseline_source='EIA WTI first 32 reported observations after 2022-05-23',
            baseline_prices_dates=[d.isoformat() for d,p in prices[split+1:split+33]],
            baseline_standardized_changes=baseline.tolist(),forecast_r=forecast_r,forecast_intercept=forecast_c))
    # External core agrees with an independently solved Bellman problem.
    forward_checks=[]
    for beta,s,cost in itertools_product([.45,.8,task_beta],NODES,[1.,2.,4.]):
        agent=make_agent(alpha,beta,cost,s);gain,_,_=solve_policy(alpha,beta,cost,s)
        error=max(abs(step(agent,p,d)-float(gain@np.array([p,d]))) for p,d in [(.3,-.7),(-.2,.8)])
        assert error<1e-11;forward_checks.append(error)
    manifest={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in EXT.iterdir() if p.is_file()}
    result=dict(commit=(EXT/'commit.txt').read_text().strip(),source_repository='https://github.com/dynamic-trading-RL/dynamic-trading',
        external_files_sha256=manifest,horizons=HORIZONS,weights=WEIGHTS.tolist(),task_clock=task_beta,
        policy_mapping='paper alpha=external risk_aversion*gamma*constant_variance; paper cost=external lam*constant_variance',
        replay_normalization='constant_variance=1; costs relative to deployment lam; risk_aversion compensated to keep paper alpha fixed',
        all_configured_matched_assets=len(asset_results),assets=asset_results,primary_eia_results=primary_results,
        repository_data_quality=data_quality,blind_cost_audits=audits,
        external_vs_bellman_checks=len(forward_checks),max_forward_error=max(forward_checks),
        limitations=['Actual historical price paths; positions are counterfactual external-code replays, not observed trades.',
                    'Quadratic coefficients are software objectives from repository settings, not measured market impact.',
                    'Positive heterogeneous risk targets are constructed at prescribed horizons; not inferred investor forecasts.',
                    'Task loss uses equal period weights and private rule losses; not investment profit or an identified aggregate loss.',
                    'Float guard is validated for these finite replays, not a general proof of source-code rounding error.'])
    (ROOT/'external_rule_and_real_path_results.json').write_text(json.dumps(result,indent=2,default=str,allow_nan=False)+'\n',encoding='utf-8')
    (ROOT/'external_paired_observations.json').write_text(json.dumps(logs,indent=2,allow_nan=False)+'\n',encoding='utf-8')
    print(json.dumps(dict(asset_count=len(asset_results),audits=audits,primary_results=primary_results,
        improvements=[(r['ticker'],r['improvement_vs_045'],r['improvement_vs_099']) for r in asset_results]),allow_nan=False))

def itertools_product(*items):
    import itertools
    return itertools.product(*items)

if __name__=='__main__':main()
